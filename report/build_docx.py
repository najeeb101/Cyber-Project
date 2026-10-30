"""Build the MS Word report from report-draft.md.

Usage (from the project root):
    python report/build_docx.py

Produces report/G02-Web-App-Security-Assessment-Report.docx: Arial 11 pt, single column,
1.0 line spacing, cover page, table of contents, page numbers, embedded figures.
The Markdown file stays the single source of truth; re-run this after every edit.
Requires python-docx (pip install python-docx).
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "report-draft.md"
OUT = HERE / "G02-Web-App-Security-Assessment-Report.docx"

FONT = "Arial"
MONO = "Consolas"
HEADING_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
MAX_IMG_WIDTH = Inches(6.2)

# Blockquotes that are notes to the team, not report content.
SKIP_QUOTES = ("**Final Markdown version", "Before submission")

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\s][^*]*\*|\[TEAM[^\]]*\])")


# --------------------------------------------------------------------------- #
# Low-level helpers
# --------------------------------------------------------------------------- #
def set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_paragraph_shading(paragraph, hex_fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    p_pr.append(shd)


def add_field(paragraph, instr: str, placeholder: str = "") -> None:
    """Insert a Word field (PAGE, TOC, ...)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    run._r.append(text)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_inline(paragraph, text: str, size: Pt | None = None, italic: bool = False) -> None:
    """Add text with **bold**, *italic*, `code` and [TEAM] highlighting."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
        elif part.startswith("**") and part.endswith("**"):
            add_inline_bold(paragraph, part[2:-2], size, italic)
            continue
        elif part.startswith("[TEAM"):
            run = paragraph.add_run(part)
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if italic:
            run.italic = True
        if size:
            run.font.size = size


def add_inline_bold(paragraph, text: str, size, italic) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
        elif part.startswith("[TEAM"):
            run = paragraph.add_run(part)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            run = paragraph.add_run(part.strip("*") if part.startswith("*") else part)
            if part.startswith("*") and part.endswith("*"):
                run.italic = True
        run.bold = True
        if italic:
            run.italic = True
        if size:
            run.font.size = size


# --------------------------------------------------------------------------- #
# Document setup
# --------------------------------------------------------------------------- #
def setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.2))

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.0
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)

    for name, size, before in (("Heading 1", 14, 14), ("Heading 2", 12, 10), ("Heading 3", 11, 8)):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = HEADING_COLOR
        rfonts = st.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)  # theme fonts would override Arial
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(2)

    # Page number footer
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE", "1")
    return doc


# --------------------------------------------------------------------------- #
# Markdown parsing
# --------------------------------------------------------------------------- #
def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def split_row(line: str) -> list[str]:
    # Cells may contain `|` only inside code spans; the report avoids that.
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc: Document, rows: list[list[str]], font_size: float = 9) -> None:
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            text = row[j] if j < len(row) else ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            add_inline(para, text, size=Pt(font_size))
            if i == 0:
                for run in para.runs:
                    run.bold = True
                set_cell_shading(cell, "D9E2F3")
                para.paragraph_format.keep_with_next = True  # never strand the header row
    header = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc: Document, code_lines: list[str]) -> None:
    for line in code_lines:
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.3)
        set_paragraph_shading(para, "F2F2F2")
        run = para.add_run(line if line else " ")
        run.font.name = MONO
        run.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc: Document, rel_path: str) -> None:
    path = (HERE / rel_path).resolve()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run()
    pic = run.add_picture(str(path))
    if pic.width > MAX_IMG_WIDTH:
        ratio = MAX_IMG_WIDTH / pic.width
        pic.width = MAX_IMG_WIDTH
        pic.height = int(pic.height * ratio)


def add_cover(doc: Document, title: str, cover_lines: list[str]) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = HEADING_COLOR
    doc.add_paragraph()

    table_lines = [l for l in cover_lines if l.startswith("|")]
    bullets = [l[2:] for l in cover_lines if l.startswith("- ") and not l.startswith("- **Title:**")]
    for b in bullets[:2]:  # course and group
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, b.replace("**Course / semester:** ", "").replace("**Group ID:** ", "Group ID: "),
                   size=Pt(13))
    doc.add_paragraph()
    add_table(doc, parse_table(table_lines), font_size=11)
    for b in bullets[2:]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, b)

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph()
    r = h.add_run("Contents")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = HEADING_COLOR
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-1" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def flush_paragraph(doc: Document, buf: list[str]) -> None:
    if buf:
        para = doc.add_paragraph()
        add_inline(para, " ".join(s.strip() for s in buf))
        buf.clear()


def build() -> None:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = setup_document()

    title = lines[0].lstrip("# ").strip()
    start_cover = next(i for i, l in enumerate(lines) if l.startswith("## 1."))
    start_body = next(i for i, l in enumerate(lines) if l.startswith("## 2."))
    add_cover(doc, title, [l for l in lines[start_cover + 1:start_body] if l.strip() and l != "---"])

    i = start_body
    buf: list[str] = []
    appendix_title_pending = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code fences
        if stripped.startswith("```"):
            flush_paragraph(doc, buf)
            j = i + 1
            while not lines[j].strip().startswith("```"):
                j += 1
            add_code(doc, lines[i + 1:j])
            i = j + 1
            continue

        # Blank line / rule ends a paragraph
        if not stripped or stripped == "---":
            flush_paragraph(doc, buf)
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            flush_paragraph(doc, buf)
            level, text = len(m.group(1)), m.group(2).strip()
            text = re.sub(r"\*\(([^)]*)\)\*", r"(\1)", text)
            if text == "Appendix (no word limit)":
                # The appendix title page break comes from Appendix A's heading below.
                appendix_title_pending = True
                i += 1
                continue
            if text.startswith("Appendix") and level <= 2:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                if appendix_title_pending:
                    doc.add_heading("Appendix", level=1)
                    appendix_title_pending = False
            heading_level = 1 if level <= 2 else 2
            doc.add_heading(text.replace("`", ""), level=heading_level)
            i += 1
            continue

        # Tables
        if stripped.startswith("|"):
            flush_paragraph(doc, buf)
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                j += 1
            add_table(doc, parse_table(lines[i:j]))
            i = j
            continue

        # Images
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", stripped)
        if m:
            flush_paragraph(doc, buf)
            add_image(doc, m.group(1))
            i += 1
            continue

        # Figure captions
        if stripped.startswith("*Figure") and stripped.endswith("*"):
            flush_paragraph(doc, buf)
            j = i
            cap = []
            while j < len(lines) and lines[j].strip():
                cap.append(lines[j].strip())
                j += 1
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(para, " ".join(cap).strip("*"), size=Pt(9), italic=True)
            i = j
            continue

        # Blockquotes
        if stripped.startswith(">"):
            flush_paragraph(doc, buf)
            j = i
            quote = []
            while j < len(lines) and lines[j].strip().startswith(">"):
                quote.append(lines[j].strip().lstrip(">").strip())
                j += 1
            text = " ".join(quote)
            if not text.startswith(SKIP_QUOTES):
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(1)
                para.paragraph_format.right_indent = Cm(1)
                add_inline(para, text, italic=True)
            i = j
            continue

        # Bullets (two-space indentation = nested)
        m = re.match(r"^(\s*)- (.*)$", line)
        if m:
            flush_paragraph(doc, buf)
            depth = len(m.group(1)) // 2
            text = m.group(2)
            j = i + 1
            while j < len(lines) and lines[j].startswith(" " * (len(m.group(1)) + 2)) \
                    and not re.match(r"^\s*(- |\d+\. )", lines[j]):
                text += " " + lines[j].strip()
                j += 1
            para = doc.add_paragraph(style="List Bullet 2" if depth else "List Bullet")
            add_inline(para, text)
            i = j
            continue

        # Numbered items (manual numbers keep the source numbering exactly)
        m = re.match(r"^(\s*)(\d+)\. (.*)$", line)
        if m:
            flush_paragraph(doc, buf)
            text = m.group(3)
            j = i + 1
            while j < len(lines) and lines[j].startswith("   ") and not re.match(r"^\s*(- |\d+\. )", lines[j]):
                text += " " + lines[j].strip()
                j += 1
            para = doc.add_paragraph()
            pf = para.paragraph_format
            pf.left_indent = Cm(0.9)
            pf.first_line_indent = Cm(-0.6)
            pf.space_after = Pt(2)
            add_inline(para, f"{m.group(2)}.\t{text}")
            i = j
            continue

        buf.append(line)
        i += 1

    flush_paragraph(doc, buf)
    doc.save(OUT)
    print(f"written: {OUT}")


if __name__ == "__main__":
    build()
